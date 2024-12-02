import streamlit as st

def add_export_button():
    """Add export buttons for different file formats"""
    import pandas as pd
    from fpdf import FPDF
    from PIL import Image
    import markdown
    import io
    import docx
    
    # Initialize session state for export buttons if not exists
    if 'show_export_buttons' not in st.session_state:
        st.session_state.show_export_buttons = False
    
    # Toggle button for showing/hiding export options
    if st.button("Export Chat" if not st.session_state.show_export_buttons else "Hide Export Options", key="export_toggle"):
        st.session_state.show_export_buttons = not st.session_state.show_export_buttons
        st.rerun()
    
    # Only show export options if state is True
    if st.session_state.show_export_buttons and 'messages' in st.session_state:
        # Convert messages to DataFrame
        df = pd.DataFrame(st.session_state.messages)
        
        # Create columns for export buttons
        col1, col2, col3 = st.columns(3)
        
        # CSV Export
        with col1:
            st.download_button(
                label="Export as CSV",
                data=df.to_csv(index=False),
                file_name='chat_history.csv',
                mime='text/csv',
                key="csv_export"
            )
        
        # Markdown Export
        with col2:
            md_content = ""
            for msg in st.session_state.messages:
                md_content += f"**{msg['role'].title()}**: {msg['content']}\n\n"
            
            st.download_button(
                label="Export as Markdown",
                data=md_content,
                file_name='chat_history.md',
                mime='text/markdown',
                key="md_export"
            )
        
        # PDF Export
        with col3:
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Arial", size=12)
            
            for msg in st.session_state.messages:
                pdf.cell(200, 10, txt=f"{msg['role'].title()}: {msg['content']}", ln=True)
            
            st.download_button(
                label="Export as PDF",
                data=pdf.output(dest='S').encode('latin-1'),
                file_name='chat_history.pdf',
                mime='application/pdf',
                key="pdf_export"
            )
        
        # Create another row of columns
        col4, col5, col6 = st.columns(3)
        
        # DOCX Export
        with col4:
            doc = docx.Document()
            for msg in st.session_state.messages:
                doc.add_paragraph(f"{msg['role'].title()}: {msg['content']}")
            
            # Save to bytes
            docx_bytes = io.BytesIO()
            doc.save(docx_bytes)
            docx_bytes.seek(0)
            
            st.download_button(
                label="Export as DOCX",
                data=docx_bytes.getvalue(),
                file_name='chat_history.docx',
                mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                key="docx_export"
            )
        
        # PNG Export
        with col5:
            # Create an image with the chat history
            img = Image.new('RGB', (800, 600), color='white')
            from PIL import ImageDraw
            d = ImageDraw.Draw(img)
            y_position = 10
            
            for msg in st.session_state.messages:
                text = f"{msg['role'].title()}: {msg['content']}"
                d.text((10, y_position), text, fill='black')
                y_position += 30
            
            # Convert to bytes
            img_bytes = io.BytesIO()
            img.save(img_bytes, format='PNG')
            img_bytes.seek(0)
            
            st.download_button(
                label="Export as PNG",
                data=img_bytes.getvalue(),
                file_name='chat_history.png',
                mime='image/png',
                key="png_export"
            )

def get_models(provider):
    if provider == "Google":
        return ["gemini-1.5-flash", "gemini-1.5-flash-8b","gemini-1.5-flash-002","gemini-1.5-pro-002", "gemini-exp-1121","gemini-exp-1114","gemini-1.5-pro-exp-0801", "gemini-1.5-pro",  "gemini-1.0-pro"]
    elif provider == "Claude":
        return ["claude-3-5-sonnet-20240620", "claude-3-opus-20240229", "claude-3-sonnet-20240229", "claude-3-haiku-20240307"]
    elif provider == "Groq":
        return [
            "llama-3.2-90b-vision-preview", "llama-3.2-11b-vision-preview", "llama-3.1-405b-reasoning", "llama-3.1-70b-versatile", "llama-3.1-8b-instant",
            "llama-3.2-1b-preview", "llama-3.2-3b-preview", 
            "llama3-groq-70b-8192-tool-use-preview", "llama3-groq-8b-8192-tool-use-preview",
            "llama3-70b-8192", "llama3-8b-8192", "mixtral-8x7b-32768", "gemma-7b-it", "gemma2-9b-it", 
            "llava-v1.5-7b-4096-preview", 
        ]
    elif provider == "OpenAI":
        return ["gpt-4", "gpt-3.5-turbo", "gpt-4o", "gpt-4o-mini"]
    elif provider == "Local(Ollama)":
        return ["local-llm"]  # You might want to populate this with available local models
    elif provider == "HuggingFace":
        return ["meta-llama/Meta-Llama-3-8B-Instruct"]  # Add more HuggingFace models as needed
    else:
        return []

def get_available_tools():
    return ["Search", "Gmail", "Finance", "Location", "Weather", "File Operations", "Shell", "Date and Time", "Media", "System", "Volume Control", "Python", "Knowledge", "Bluetooth", "WhatsApp", "Alarm", "Screenshare", "Note Taking", "To-Do List"]

def generate_response(user_input, session_state):
    import streamlit as st
    from langchain_community.callbacks.streamlit.streamlit_callback_handler import StreamlitCallbackHandler
    st_callback = StreamlitCallbackHandler(st.container())

    if session_state.use_agent and session_state.agent:
        response = session_state.agent.invoke({"input": user_input}, {"callbacks": [st_callback]})
        answer = response['output']
    else:
        response = session_state.chain.invoke(user_input)
        answer = response['response']
    return answer

def log_chat(user_input, answer):
    with open('chats.txt', 'a') as f:
        f.write(f'\nHuman: {user_input}\n')
        f.write(f'AI: {answer}\n')